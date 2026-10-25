import io
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import fitz
import pytest
from app import app

def create_sample_pdf(text="Hello World PaperPilot Test PDF", pages=2):
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text(fitz.Point(50, 50), f"{text} - Page {i + 1}", fontsize=14)
    stream = io.BytesIO()
    doc.save(stream)
    doc.close()
    stream.seek(0)
    return stream

@pytest.fixture
def client():
    app.config['TESTING'] = True
    with app.test_client() as client:
        yield client

def test_health(client):
    res = client.get('/api/health')
    assert res.status_code == 200
    data = res.get_json()
    assert data['status'] == 'healthy'
    assert 'PaperPilot' in data['service']

def test_merge_pdfs(client):
    pdf1 = create_sample_pdf("Doc A", 1)
    pdf2 = create_sample_pdf("Doc B", 1)

    data = {
        'files': [
            (pdf1, 'doc_a.pdf'),
            (pdf2, 'doc_b.pdf')
        ]
    }
    res = client.post('/api/pdf/merge', data=data, content_type='multipart/form-data')
    assert res.status_code == 200
    res_data = res.get_json()
    assert res_data['success'] is True
    assert 'download_url' in res_data

    # Verify download works
    dl_res = client.get(res_data['download_url'])
    assert dl_res.status_code == 200
    merged_doc = fitz.open(stream=dl_res.data, filetype="pdf")
    assert len(merged_doc) == 2
    merged_doc.close()

def test_split_pdf(client):
    pdf = create_sample_pdf("Multi Page", 3)
    data = {
        'file': (pdf, 'multi.pdf'),
        'mode': 'all'
    }
    res = client.post('/api/pdf/split', data=data, content_type='multipart/form-data')
    assert res.status_code == 200
    res_data = res.get_json()
    assert res_data['success'] is True
    assert res_data['page_count'] == 3

def test_compress_pdf(client):
    pdf = create_sample_pdf("Compress Test", 3)
    data = {
        'file': (pdf, 'test_comp.pdf'),
        'level': 'medium'
    }
    res = client.post('/api/pdf/compress', data=data, content_type='multipart/form-data')
    assert res.status_code == 200
    res_data = res.get_json()
    assert res_data['success'] is True
    assert 'compressed_size' in res_data

def test_rotate_pdf(client):
    pdf = create_sample_pdf("Rotate Test", 1)
    data = {
        'file': (pdf, 'rotate.pdf'),
        'rotation': '90',
        'pages': 'all'
    }
    res = client.post('/api/pdf/rotate', data=data, content_type='multipart/form-data')
    assert res.status_code == 200
    res_data = res.get_json()
    assert res_data['success'] is True

def test_protect_and_unlock(client):
    pdf = create_sample_pdf("Top Secret", 1)
    data = {
        'file': (pdf, 'secret.pdf'),
        'password': 'Password123!'
    }
    prot_res = client.post('/api/pdf/protect', data=data, content_type='multipart/form-data')
    assert prot_res.status_code == 200
    prot_data = prot_res.get_json()
    assert prot_data['success'] is True

    # Download protected PDF
    dl_res = client.get(prot_data['download_url'])
    assert dl_res.status_code == 200

    # Try unlocking it
    unlock_data = {
        'file': (io.BytesIO(dl_res.data), 'secret_locked.pdf'),
        'password': 'Password123!'
    }
    unlock_res = client.post('/api/pdf/unlock', data=unlock_data, content_type='multipart/form-data')
    assert unlock_res.status_code == 200
    assert unlock_res.get_json()['success'] is True

def test_watermark_pdf(client):
    pdf = create_sample_pdf("Watermark Test", 1)
    data = {
        'file': (pdf, 'wm.pdf'),
        'text': 'STRICTLY CONFIDENTIAL'
    }
    res = client.post('/api/pdf/watermark', data=data, content_type='multipart/form-data')
    assert res.status_code == 200
    assert res.get_json()['success'] is True

def test_summarize_pdf(client):
    pdf = create_sample_pdf("This is a strategic review document. The team must review the quarterly objectives. All deliverables should be approved before the final deadline.", 1)
    data = {
        'file': (pdf, 'review.pdf')
    }
    res = client.post('/api/pdf/summarize', data=data, content_type='multipart/form-data')
    assert res.status_code == 200
    res_data = res.get_json()
    assert res_data['success'] is True
    assert 'analysis' in res_data

def test_convert_markdown(client):
    pdf = create_sample_pdf("Document Title\nThis is sample markdown paragraph.", 1)
    data = {
        'file': (pdf, 'doc.pdf')
    }
    res = client.post('/api/pdf/convert/markdown', data=data, content_type='multipart/form-data')
    assert res.status_code == 200
    res_data = res.get_json()
    assert res_data['success'] is True
    assert res_data['filename'] == 'doc.md'

def test_crop_pdf(client):
    pdf = create_sample_pdf("Crop Content", 1)
    data = {
        'file': (pdf, 'crop_test.pdf'),
        'top': 10,
        'bottom': 10,
        'left': 10,
        'right': 10
    }
    res = client.post('/api/pdf/crop', data=data, content_type='multipart/form-data')
    assert res.status_code == 200
    res_data = res.get_json()
    assert res_data['success'] is True

def test_page_numbers(client):
    pdf = create_sample_pdf("Numbered Document", 2)
    data = {
        'file': (pdf, 'numbered.pdf'),
        'position': 'bottom-right',
        'start_num': 1
    }
    res = client.post('/api/pdf/page-numbers', data=data, content_type='multipart/form-data')
    assert res.status_code == 200
    res_data = res.get_json()
    assert res_data['success'] is True

def test_word_to_pdf(client):
    import docx
    doc = docx.Document()
    doc.add_heading('PaperPilot Word Conversion Test', level=1)
    doc.add_paragraph('This is a test document converted to PDF.')
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    res = client.post('/api/pdf/convert/word-to-pdf', data={'file': (stream, 'test.docx')}, content_type='multipart/form-data')
    assert res.status_code == 200
    assert res.get_json()['success'] is True

def test_excel_to_pdf(client):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["ID", "Name", "Score"])
    ws.append([1, "Alpha", 95])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    res = client.post('/api/pdf/convert/excel-to-pdf', data={'file': (stream, 'test.xlsx')}, content_type='multipart/form-data')
    assert res.status_code == 200
    assert res.get_json()['success'] is True

def test_powerpoint_to_pdf(client):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    stream = io.BytesIO()
    prs.save(stream)
    stream.seek(0)

    res = client.post('/api/pdf/convert/powerpoint-to-pdf', data={'file': (stream, 'test.pptx')}, content_type='multipart/form-data')
    assert res.status_code == 200
    assert res.get_json()['success'] is True

def test_pdf_to_pdfa(client):
    pdf = create_sample_pdf("Archival Test", 1)
    res = client.post('/api/pdf/convert/pdfa', data={'file': (pdf, 'archive.pdf')}, content_type='multipart/form-data')
    assert res.status_code == 200
    res_data = res.get_json()
    assert res_data['success'] is True
    assert 'conformance' in res_data

def test_scan_to_pdf(client):
    from PIL import Image
    img = Image.new('RGB', (100, 100), color='white')
    stream = io.BytesIO()
    img.save(stream, format='JPEG')
    stream.seek(0)

    res = client.post('/api/pdf/convert/scan-to-pdf', data={'files': [(stream, 'scan.jpg')]}, content_type='multipart/form-data')
    assert res.status_code == 200
    assert res.get_json()['success'] is True

def test_sign_pdf(client):
    pdf = create_sample_pdf("Document For Signature", 1)
    # 1x1 transparent PNG data URI
    dummy_sig = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
    res = client.post('/api/pdf/sign', data={
        'file': (pdf, 'contract.pdf'),
        'signature_data': dummy_sig,
        'page': 1,
        'x': 50,
        'y': 50
    }, content_type='multipart/form-data')
    assert res.status_code == 200
    assert res.get_json()['success'] is True

def test_edit_pdf(client):
    pdf = create_sample_pdf("Editable PDF", 1)
    res = client.post('/api/pdf/edit', data={
        'file': (pdf, 'edit_me.pdf'),
        'text': 'Approved by PaperPilot',
        'page': 1,
        'x': 50,
        'y': 100
    }, content_type='multipart/form-data')
    assert res.status_code == 200
    assert res.get_json()['success'] is True

def test_forms_detect_and_fill(client):
    pdf1 = create_sample_pdf("Forms PDF", 1)
    detect_res = client.post('/api/pdf/forms/detect', data={'file': (pdf1, 'form.pdf')}, content_type='multipart/form-data')
    assert detect_res.status_code == 200
    assert detect_res.get_json()['success'] is True

    pdf2 = create_sample_pdf("Forms PDF", 1)
    fill_res = client.post('/api/pdf/forms/fill', data={
        'file': (pdf2, 'form.pdf'),
        'fields': '{"username": "John Doe"}'
    }, content_type='multipart/form-data')
    assert fill_res.status_code == 200
    assert fill_res.get_json()['success'] is True

def test_translate_pdf(client):
    pdf = create_sample_pdf("This document is about flight and document navigation under control.", 1)
    res = client.post('/api/pdf/translate', data={
        'file': (pdf, 'travel.pdf'),
        'target_lang': 'es'
    }, content_type='multipart/form-data')
    assert res.status_code == 200
    res_data = res.get_json()
    assert res_data['success'] is True
    assert 'translation' in res_data

